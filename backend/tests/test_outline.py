import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document as DocxDocument

from app.ingestion.outline import OutlineEntry, extract_outline


def _make_docx(sections):
    """sections: list of (heading, [paragraph, ...])"""
    doc = DocxDocument()
    for heading, paragraphs in sections:
        if heading:
            doc.add_heading(heading, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    # Windows không cho ghi đè khi handle còn mở (WinError 32)
    tmp.close()
    doc.save(tmp.name)
    return tmp.name


class TestExtractOutlineDocx(unittest.TestCase):
    def test_extracts_headings_in_order(self):
        path = _make_docx(
            [
                ("Gradient Descent", ["Nội dung một.", "Nội dung hai."]),
                ("Learning Rate", ["Nội dung ba."]),
            ]
        )
        try:
            outline = extract_outline(path)
            self.assertEqual([e.title for e in outline], ["Gradient Descent", "Learning Rate"])
            self.assertEqual([e.order for e in outline], [0, 1])
        finally:
            os.remove(path)

    def test_entries_carry_position_ref_matching_parser(self):
        # parser gom 10 đoạn/section, nên heading đầu tiên phải nằm ở "Mục 1"
        path = _make_docx([("Chương đầu", ["Một câu."])])
        try:
            outline = extract_outline(path)
            self.assertEqual(outline[0].position_ref, "Mục 1")
        finally:
            os.remove(path)

    def test_heading_in_later_section_gets_later_position_ref(self):
        # 12 đoạn văn trước heading thứ hai -> heading đó rơi sang Mục 2
        first = ("Chương một", [f"Đoạn số {i}." for i in range(12)])
        second = ("Chương hai", ["Nội dung chương hai."])
        path = _make_docx([first, second])
        try:
            outline = extract_outline(path)
            self.assertEqual(outline[0].position_ref, "Mục 1")
            self.assertEqual(outline[1].position_ref, "Mục 2")
        finally:
            os.remove(path)

    def test_document_without_headings_returns_empty(self):
        path = _make_docx([(None, ["Chỉ có đoạn văn thường.", "Không heading nào."])])
        try:
            self.assertEqual(extract_outline(path), [])
        finally:
            os.remove(path)

    def test_returns_outline_entry_instances(self):
        path = _make_docx([("Tiêu đề", ["Nội dung."])])
        try:
            self.assertIsInstance(extract_outline(path)[0], OutlineEntry)
        finally:
            os.remove(path)

    def test_duplicate_headings_are_kept_once(self):
        path = _make_docx([("Trùng", ["A."]), ("Trùng", ["B."])])
        try:
            outline = extract_outline(path)
            self.assertEqual([e.title for e in outline], ["Trùng"])
        finally:
            os.remove(path)

    def test_unsupported_extension_returns_empty(self):
        tmp = tempfile.NamedTemporaryFile(suffix=".txt", delete=False)
        tmp.close()
        try:
            self.assertEqual(extract_outline(tmp.name), [])
        finally:
            os.remove(tmp.name)


if __name__ == "__main__":
    unittest.main()
