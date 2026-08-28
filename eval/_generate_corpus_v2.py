"""Mở rộng corpus đánh giá từ 4 lên 15 tài liệu (Golden Set v2).

Chạy MỘT LẦN để tạo file trong eval/documents/. Không phải một phần của bộ
test tự động của backend.

Bốn tài liệu gốc (ML_Optimization, ML_DecisionTree, DL_CNN, DL_NeuralNetwork)
do eval/_generate_corpus.py sinh ra và ĐƯỢC GIỮ NGUYÊN — script này chỉ thêm
11 tài liệu mới.

Ba nhóm mới được thêm có chủ đích:

1. Ba tài liệu PDF. Corpus cũ toàn DOCX nên đường phân tích PDF
   (app/ingestion/parser.py) và tính năng mở tài liệu gốc đúng trang chưa hề
   được đánh giá lần nào, dù PDF là một trong hai định dạng PRD cam kết hỗ trợ.

2. Ba tài liệu tiếng Anh. Cần để đo yêu cầu "trả lời đúng ngôn ngữ của câu
   hỏi" trên tài liệu thật, thay vì chỉ suy đoán.

3. Hai tài liệu DÀI cố ý chôn một thuật ngữ ở sâu, xuất hiện đúng một lần.
   Đây là bài kiểm tra trực tiếp cho cơ chế truy hồi hai lượt: truy hồi ngữ
   nghĩa hay bỏ sót đúng dạng này, và lượt mở rộng theo từ khoá phải cứu được.

Cần: pip install reportlab (chỉ dùng cho việc sinh corpus, KHÔNG phải
dependency lúc chạy ứng dụng nên không khai trong requirements.txt).
"""

import os

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

OUT_DIR = os.path.join(os.path.dirname(__file__), "documents")
os.makedirs(OUT_DIR, exist_ok=True)


def write_docx(filename, sections):
    doc = Document()
    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)
    doc.save(os.path.join(OUT_DIR, filename))
    print("Wrote", filename)


def write_pdf(filename, sections, page_break_every=2):
    """Mỗi `page_break_every` section thì sang trang mới, để tài liệu có nhiều
    trang thật — cần cho việc kiểm tra citation trỏ đúng số trang."""
    doc = SimpleDocTemplate(os.path.join(OUT_DIR, filename), pagesize=A4)
    styles = getSampleStyleSheet()
    flow = []
    for i, (heading, paragraphs) in enumerate(sections):
        flow.append(Paragraph(heading, styles["Heading1"]))
        for p in paragraphs:
            flow.append(Paragraph(p, styles["BodyText"]))
            flow.append(Spacer(1, 6))
        if (i + 1) % page_break_every == 0 and i + 1 < len(sections):
            flow.append(PageBreak())
    doc.build(flow)
    print("Wrote", filename)


# ============================================================
# Nhóm 1 — Machine Learning tiếng Việt
# ============================================================

write_docx(
    "ML_Regression.docx",
    [
        (
            "Hồi quy tuyến tính",
            [
                "Hồi quy tuyến tính là mô hình dự đoán một giá trị liên tục bằng cách giả định "
                "quan hệ tuyến tính giữa biến đầu vào và biến đầu ra. Mô hình có dạng "
                "y = w1*x1 + w2*x2 + ... + b, trong đó các trọng số w và hệ số chệch b được học "
                "từ dữ liệu huấn luyện.",
                "Hàm mất mát thường dùng là sai số bình phương trung bình (Mean Squared Error), "
                "tính trung bình bình phương chênh lệch giữa giá trị dự đoán và giá trị thật. "
                "Bình phương khiến sai số lớn bị phạt nặng hơn sai số nhỏ, nên mô hình nhạy cảm "
                "với điểm ngoại lai.",
            ],
        ),
        (
            "Hồi quy logistic",
            [
                "Hồi quy logistic dùng cho bài toán phân loại nhị phân, không phải dự đoán giá "
                "trị liên tục. Mô hình tính một tổ hợp tuyến tính của đầu vào rồi đưa qua hàm "
                "sigmoid để ép kết quả về khoảng từ 0 đến 1, diễn giải được như xác suất thuộc "
                "lớp dương.",
                "Hàm mất mát dùng ở đây là entropy chéo (cross-entropy) chứ không phải sai số "
                "bình phương, vì cross-entropy phạt rất nặng những dự đoán tự tin nhưng sai, và "
                "cho gradient ổn định hơn khi kết hợp với sigmoid.",
            ],
        ),
        (
            "Đánh giá mô hình hồi quy",
            [
                "Hệ số xác định R bình phương cho biết tỉ lệ phương sai của biến mục tiêu được "
                "mô hình giải thích. Giá trị càng gần 1 thì mô hình càng khớp dữ liệu, nhưng R "
                "bình phương luôn tăng khi thêm biến, kể cả biến vô nghĩa, nên cần dùng R bình "
                "phương hiệu chỉnh khi so sánh các mô hình có số biến khác nhau.",
                "Sai số tuyệt đối trung bình (Mean Absolute Error) là một lựa chọn thay thế cho "
                "MSE, ít nhạy cảm với điểm ngoại lai hơn vì không bình phương sai số.",
            ],
        ),
    ],
)

write_docx(
    "ML_Overfitting.docx",
    [
        (
            "Quá khớp và dưới khớp",
            [
                "Quá khớp (overfitting) xảy ra khi mô hình học thuộc cả nhiễu trong dữ liệu "
                "huấn luyện, dẫn tới sai số huấn luyện rất thấp nhưng sai số trên dữ liệu mới "
                "lại cao. Dưới khớp (underfitting) là tình trạng ngược lại: mô hình quá đơn "
                "giản nên không nắm được quy luật, sai số cao ở cả hai tập.",
                "Dấu hiệu nhận biết quá khớp rõ nhất là khoảng cách ngày càng rộng giữa đường "
                "sai số huấn luyện và đường sai số kiểm định khi tăng số vòng huấn luyện.",
            ],
        ),
        (
            "Điều chuẩn",
            [
                "Điều chuẩn L2 (ridge) cộng thêm tổng bình phương các trọng số vào hàm mất mát, "
                "khiến mô hình ưu tiên các trọng số nhỏ và trải đều thay vì dồn vào một vài "
                "biến. Điều chuẩn L1 (lasso) cộng tổng trị tuyệt đối của trọng số, có xu hướng "
                "đẩy hẳn một số trọng số về 0 nên đồng thời làm nhiệm vụ chọn biến.",
                "Dropout là kỹ thuật điều chuẩn dành riêng cho mạng nơ-ron: trong mỗi vòng "
                "huấn luyện, một tỉ lệ nơ-ron được tắt ngẫu nhiên, buộc mạng không phụ thuộc "
                "quá mức vào bất kỳ nơ-ron đơn lẻ nào.",
            ],
        ),
        (
            "Kiểm định chéo",
            [
                "Kiểm định chéo k-fold chia dữ liệu thành k phần bằng nhau, lần lượt lấy một "
                "phần làm tập kiểm định và k-1 phần còn lại làm tập huấn luyện, rồi lấy trung "
                "bình kết quả của k lần chạy. Cách này cho ước lượng hiệu năng ổn định hơn so "
                "với chỉ chia một lần.",
                "Với dữ liệu mất cân bằng lớp, nên dùng kiểm định chéo phân tầng (stratified) "
                "để mỗi fold giữ nguyên tỉ lệ các lớp như tập gốc.",
            ],
        ),
    ],
)

write_docx(
    "DATA_Preprocessing.docx",
    [
        (
            "Chuẩn hoá dữ liệu",
            [
                "Chuẩn hoá min-max đưa mọi giá trị của một đặc trưng về khoảng từ 0 đến 1 bằng "
                "công thức (x - min) / (max - min). Cách này giữ nguyên hình dạng phân phối "
                "nhưng rất nhạy với điểm ngoại lai, vì một giá trị cực đại bất thường sẽ nén "
                "toàn bộ phần còn lại về gần 0.",
                "Chuẩn hoá z-score biến đổi dữ liệu về trung bình 0 và độ lệch chuẩn 1 theo "
                "công thức (x - trung bình) / độ lệch chuẩn. Cách này ít nhạy với ngoại lai hơn "
                "min-max và là lựa chọn mặc định cho phần lớn thuật toán dựa trên khoảng cách.",
            ],
        ),
        (
            "Xử lý giá trị thiếu",
            [
                "Xoá hẳn dòng có giá trị thiếu là cách đơn giản nhất nhưng làm mất dữ liệu, chỉ "
                "nên dùng khi tỉ lệ thiếu rất nhỏ. Điền bằng giá trị trung bình hoặc trung vị "
                "giữ được kích thước tập dữ liệu nhưng làm giảm phương sai thật của đặc trưng.",
                "Với dữ liệu chuỗi thời gian, điền tiến (forward fill) lấy giá trị quan sát gần "
                "nhất trước đó thường hợp lý hơn điền bằng trung bình, vì nó tôn trọng thứ tự "
                "thời gian.",
            ],
        ),
        (
            "Mã hoá biến hạng mục",
            [
                "Mã hoá one-hot tạo một cột nhị phân cho mỗi giá trị hạng mục, tránh áp đặt thứ "
                "tự giả lên dữ liệu không có thứ tự. Nhược điểm là số chiều bùng nổ khi biến có "
                "quá nhiều giá trị khác nhau.",
                "Mã hoá thứ tự (ordinal) gán mỗi hạng mục một số nguyên, chỉ phù hợp khi các "
                "hạng mục thật sự có thứ tự tự nhiên như thấp, trung bình, cao.",
            ],
        ),
    ],
)

write_docx(
    "STATS_Probability.docx",
    [
        (
            "Xác suất có điều kiện và định lý Bayes",
            [
                "Xác suất có điều kiện P(A|B) là xác suất xảy ra A khi biết B đã xảy ra, tính "
                "bằng P(A giao B) chia cho P(B). Định lý Bayes cho phép đảo chiều điều kiện: "
                "P(A|B) = P(B|A) * P(A) / P(B).",
                "Trong phân loại, P(A) gọi là xác suất tiên nghiệm, P(B|A) là khả năng "
                "(likelihood), và P(A|B) là xác suất hậu nghiệm. Bộ phân loại Naive Bayes áp "
                "dụng công thức này kèm giả định các đặc trưng độc lập với nhau khi đã biết lớp.",
            ],
        ),
        (
            "Phân phối thường gặp",
            [
                "Phân phối chuẩn có dạng hình chuông đối xứng, xác định bởi hai tham số là "
                "trung bình và độ lệch chuẩn. Khoảng một trăm phần trăm dữ liệu nằm trong ba "
                "độ lệch chuẩn quanh trung bình, cụ thể khoảng 68 phần trăm trong một độ lệch "
                "chuẩn và khoảng 95 phần trăm trong hai độ lệch chuẩn.",
                "Phân phối nhị thức mô tả số lần thành công trong n phép thử độc lập với xác "
                "suất thành công cố định. Khi n đủ lớn, phân phối nhị thức xấp xỉ phân phối "
                "chuẩn.",
            ],
        ),
        (
            "Phương sai và độ chệch",
            [
                "Sai số dự đoán của một mô hình phân tách được thành ba phần: độ chệch bình "
                "phương, phương sai, và nhiễu không thể giảm. Độ chệch cao tương ứng với mô "
                "hình quá đơn giản, phương sai cao tương ứng với mô hình quá nhạy với dữ liệu "
                "huấn luyện cụ thể.",
                "Đánh đổi độ chệch và phương sai là lý do không thể giảm cả hai cùng lúc bằng "
                "cách chỉ thay đổi độ phức tạp mô hình; muốn giảm cả hai thường phải thêm dữ "
                "liệu hoặc thêm đặc trưng tốt hơn.",
            ],
        ),
    ],
)

# ============================================================
# Nhóm 2 — Deep Learning tiếng Việt
# ============================================================

write_docx(
    "DL_RNN_LSTM.docx",
    [
        (
            "Mạng nơ-ron hồi quy",
            [
                "Mạng nơ-ron hồi quy (RNN) xử lý dữ liệu tuần tự bằng cách duy trì một trạng "
                "thái ẩn được cập nhật ở mỗi bước thời gian, cho phép thông tin từ các bước "
                "trước ảnh hưởng tới đầu ra hiện tại. Cùng một bộ trọng số được dùng lại ở mọi "
                "bước thời gian.",
                "Điểm yếu cố hữu của RNN cơ bản là khó học được phụ thuộc xa: khi lan truyền "
                "ngược qua nhiều bước thời gian, gradient bị nhân liên tiếp với các giá trị nhỏ "
                "hơn 1 và tiêu biến dần.",
            ],
        ),
        (
            "Tiêu biến gradient",
            [
                "Hiện tượng tiêu biến gradient xảy ra khi đạo hàm nhân dồn qua nhiều lớp hoặc "
                "nhiều bước thời gian trở nên cực nhỏ, khiến các lớp đầu gần như không được cập "
                "nhật. Hàm kích hoạt sigmoid và tanh làm hiện tượng này nặng hơn vì đạo hàm của "
                "chúng luôn nhỏ hơn 1.",
                "Các cách giảm nhẹ gồm dùng hàm kích hoạt ReLU, khởi tạo trọng số hợp lý, thêm "
                "kết nối tắt như trong ResNet, và với dữ liệu tuần tự thì dùng kiến trúc có cổng "
                "như LSTM.",
            ],
        ),
        (
            "LSTM và cơ chế cổng",
            [
                "LSTM giải quyết vấn đề phụ thuộc xa bằng một ô nhớ (cell state) chạy xuyên "
                "suốt chuỗi, cùng ba cổng điều khiển: cổng quên quyết định bỏ bớt thông tin cũ, "
                "cổng vào quyết định ghi thêm thông tin mới, cổng ra quyết định phần nào của ô "
                "nhớ được đưa ra ngoài.",
                "Nhờ ô nhớ được cập nhật bằng phép cộng thay vì nhân liên tiếp, gradient chảy "
                "qua nhiều bước thời gian mà không tiêu biến nhanh như RNN cơ bản.",
            ],
        ),
    ],
)

write_docx(
    "DL_Transformer.docx",
    [
        (
            "Cơ chế chú ý",
            [
                "Cơ chế chú ý (attention) cho phép mô hình, khi xử lý một vị trí trong chuỗi, "
                "nhìn trực tiếp tới mọi vị trí khác và gán trọng số cho từng vị trí theo mức độ "
                "liên quan. Mỗi vị trí sinh ra ba vector: truy vấn (query), khoá (key) và giá "
                "trị (value).",
                "Trọng số chú ý được tính bằng tích vô hướng giữa truy vấn và khoá, chia cho căn "
                "bậc hai của số chiều khoá rồi đưa qua softmax. Phép chia này giữ cho tích vô "
                "hướng không quá lớn khi số chiều tăng, tránh softmax bão hoà.",
            ],
        ),
        (
            "Kiến trúc Transformer",
            [
                "Transformer bỏ hẳn tính hồi quy, chỉ dùng chú ý và mạng truyền thẳng, nhờ đó "
                "toàn bộ chuỗi được xử lý song song thay vì tuần tự từng bước như RNN. Đây là "
                "lý do chính khiến Transformer huấn luyện nhanh hơn nhiều trên phần cứng song "
                "song.",
                "Vì không còn tính tuần tự, mô hình mất thông tin về thứ tự nên phải cộng thêm "
                "mã hoá vị trí (positional encoding) vào biểu diễn đầu vào.",
            ],
        ),
        (
            "Chú ý đa đầu",
            [
                "Chú ý đa đầu chạy nhiều cơ chế chú ý song song trên các không gian con khác "
                "nhau của biểu diễn, rồi ghép kết quả lại. Mỗi đầu có thể học một kiểu quan hệ "
                "riêng, ví dụ một đầu bắt quan hệ cú pháp gần, một đầu bắt quan hệ ngữ nghĩa xa.",
                "Số đầu và số chiều mỗi đầu thường được chọn sao cho tích của chúng bằng tổng "
                "số chiều của mô hình, nên chi phí tính toán không tăng so với một đầu duy nhất "
                "có cùng tổng số chiều.",
            ],
        ),
    ],
)

# ============================================================
# Nhóm 3 — Tài liệu tiếng Anh
# ============================================================

write_docx(
    "EN_Ensemble_Methods.docx",
    [
        (
            "Bagging",
            [
                "Bagging, short for bootstrap aggregating, trains many models independently on "
                "different bootstrap samples of the training data and averages their predictions. "
                "Because the models are trained independently, bagging mainly reduces variance "
                "rather than bias.",
                "Random Forest is bagging applied to decision trees, with one extra source of "
                "randomness: at every split, only a random subset of features is considered. "
                "This decorrelates the individual trees and improves the ensemble further.",
            ],
        ),
        (
            "Boosting",
            [
                "Boosting trains models sequentially, where each new model focuses on the "
                "examples that previous models got wrong. Unlike bagging, the models are not "
                "independent, and boosting reduces bias as well as variance.",
                "Gradient boosting frames this as gradient descent in function space: each new "
                "tree is fitted to the negative gradient of the loss with respect to the current "
                "ensemble prediction, which for squared error is simply the residual.",
            ],
        ),
        (
            "Choosing between them",
            [
                "Bagging is easy to parallelise and is quite robust to hyperparameter choices, "
                "which makes it a safe default. Boosting usually reaches higher accuracy but is "
                "sequential, slower to train, and more sensitive to the learning rate and the "
                "number of estimators.",
                "Boosting is also more prone to overfitting noisy data, because it keeps "
                "increasing the weight of examples it cannot fit, including mislabelled ones.",
            ],
        ),
    ],
)

write_pdf(
    "EN_Reinforcement_Learning.pdf",
    [
        (
            "Agents and Environments",
            [
                "Reinforcement learning studies an agent that interacts with an environment over "
                "discrete time steps. At each step the agent observes a state, chooses an action, "
                "and receives a scalar reward together with the next state.",
                "The agent's goal is to maximise the expected cumulative reward, not the "
                "immediate reward. Future rewards are usually discounted by a factor between 0 "
                "and 1, so that rewards further in the future count less.",
            ],
        ),
        (
            "Value Functions",
            [
                "The state value function gives the expected return when starting from a state "
                "and following a given policy thereafter. The action value function, often "
                "written Q, gives the expected return when taking a specific action in a state "
                "and following the policy afterwards.",
                "The Bellman equation expresses the value of a state in terms of the values of "
                "its successor states, which is what makes dynamic programming and temporal "
                "difference learning possible.",
            ],
        ),
        (
            "Exploration and Exploitation",
            [
                "An agent that always takes the action it currently believes is best may never "
                "discover a better one. An agent that always explores randomly never accumulates "
                "reward. Balancing the two is the exploration-exploitation trade-off.",
                "The epsilon-greedy strategy is the simplest practical answer: with probability "
                "epsilon take a random action, otherwise take the action with the highest "
                "estimated value. Epsilon is usually decayed over training.",
            ],
        ),
        (
            "Q-Learning",
            [
                "Q-learning is an off-policy temporal difference method that updates the action "
                "value estimate towards the reward plus the discounted maximum value of the next "
                "state. Being off-policy means it learns about the greedy policy while behaving "
                "according to a different, more exploratory policy.",
                "Q-learning converges to the optimal action value function under the conditions "
                "that every state-action pair is visited infinitely often and the learning rate "
                "decays appropriately.",
            ],
        ),
    ],
)

write_pdf(
    "EN_Model_Evaluation.pdf",
    [
        (
            "Confusion Matrix",
            [
                "A confusion matrix summarises classification results as counts of true "
                "positives, false positives, true negatives, and false negatives. Every other "
                "classification metric is derived from these four numbers.",
                "Accuracy is the proportion of correct predictions overall. It is misleading on "
                "imbalanced datasets: a classifier that always predicts the majority class can "
                "reach very high accuracy while being useless.",
            ],
        ),
        (
            "Precision, Recall and F1",
            [
                "Precision is the fraction of predicted positives that are actually positive, "
                "so it answers the question of how much we can trust a positive prediction. "
                "Recall is the fraction of actual positives that the model found, so it answers "
                "how much of the target class we are missing.",
                "The F1 score is the harmonic mean of precision and recall. The harmonic mean is "
                "used rather than the arithmetic mean because it punishes a large imbalance "
                "between the two values.",
            ],
        ),
        (
            "ROC and AUC",
            [
                "The ROC curve plots the true positive rate against the false positive rate as "
                "the classification threshold is varied. A model that ranks all positives above "
                "all negatives traces a curve through the top-left corner.",
                "The area under the ROC curve, or AUC, summarises the curve as a single number "
                "and can be interpreted as the probability that a randomly chosen positive is "
                "ranked above a randomly chosen negative.",
            ],
        ),
    ],
)

# ============================================================
# Nhóm 4 — Tài liệu DÀI, cố ý chôn thuật ngữ ở sâu
# ============================================================

# "Kernel Trick" chỉ xuất hiện ĐÚNG MỘT LẦN, ở section thứ 8/9. Truy hồi ngữ
# nghĩa dễ bỏ sót vì phần lớn tài liệu nói về chuyện khác; lượt truy hồi mở
# rộng theo từ khoá phải bắt được.
_HANDBOOK_FILLER = [
    (
        "Chương 1 — Tổng quan quy trình học máy",
        [
            "Một dự án học máy điển hình đi qua các bước: xác định bài toán, thu thập dữ liệu, "
            "khám phá và làm sạch dữ liệu, chọn và huấn luyện mô hình, đánh giá, rồi triển khai "
            "và theo dõi. Các bước này hiếm khi tuyến tính; kết quả đánh giá thường buộc quay "
            "lại bước dữ liệu.",
            "Phần lớn thời gian thực tế của một dự án nằm ở bước chuẩn bị dữ liệu chứ không "
            "phải bước chọn mô hình.",
        ],
    ),
    (
        "Chương 2 — Thu thập và gán nhãn dữ liệu",
        [
            "Chất lượng nhãn quyết định trần hiệu năng của mô hình giám sát. Khi nhiều người "
            "cùng gán nhãn, cần đo mức độ đồng thuận giữa các người gán để phát hiện hướng dẫn "
            "gán nhãn mơ hồ.",
            "Dữ liệu thu thập từ một nguồn duy nhất thường mang thiên lệch của nguồn đó, và mô "
            "hình sẽ học luôn thiên lệch ấy.",
        ],
    ),
    (
        "Chương 3 — Khám phá dữ liệu",
        [
            "Thống kê mô tả và biểu đồ phân phối giúp phát hiện sớm giá trị bất thường, phân "
            "phối lệch, và các cột gần như hằng số không mang thông tin.",
            "Ma trận tương quan chỉ ra quan hệ tuyến tính giữa các đặc trưng, hữu ích để phát "
            "hiện đa cộng tuyến trước khi huấn luyện mô hình tuyến tính.",
        ],
    ),
    (
        "Chương 4 — Kỹ thuật đặc trưng",
        [
            "Tạo đặc trưng mới từ đặc trưng có sẵn thường mang lại cải thiện lớn hơn việc đổi "
            "sang một thuật toán phức tạp hơn. Ví dụ tách ngày tháng thành thứ trong tuần, "
            "tháng, và cờ ngày lễ.",
            "Cần cẩn thận tránh rò rỉ dữ liệu: mọi phép biến đổi có dùng thống kê toàn cục phải "
            "được học trên tập huấn luyện rồi mới áp dụng cho tập kiểm định.",
        ],
    ),
    (
        "Chương 5 — Lựa chọn mô hình",
        [
            "Nên bắt đầu bằng một mô hình cơ sở đơn giản để có mốc so sánh. Nếu một mô hình phức "
            "tạp không vượt được mốc này đáng kể thì độ phức tạp thêm vào là không đáng.",
            "Tìm kiếm siêu tham số theo lưới đầy đủ tốn kém khi số siêu tham số lớn; tìm kiếm "
            "ngẫu nhiên thường đạt kết quả tương đương với ít lần thử hơn nhiều.",
        ],
    ),
    (
        "Chương 6 — Huấn luyện và theo dõi",
        [
            "Đường cong học tập theo số vòng huấn luyện cho biết mô hình đang dưới khớp hay quá "
            "khớp, và khi nào nên dừng sớm.",
            "Dừng sớm dựa trên sai số của tập kiểm định là một dạng điều chuẩn ngầm, rẻ và hiệu "
            "quả.",
        ],
    ),
    (
        "Chương 7 — Triển khai",
        [
            "Mô hình đưa vào vận hành cần được theo dõi trôi dạt dữ liệu: phân phối đầu vào thực "
            "tế có thể lệch dần khỏi phân phối lúc huấn luyện, làm hiệu năng giảm mà không có "
            "lỗi kỹ thuật nào.",
            "Nên giữ khả năng quay lại phiên bản mô hình trước đó, vì hồi quy chất lượng thường "
            "chỉ lộ ra sau khi đã phục vụ người dùng thật.",
        ],
    ),
    (
        "Chương 8 — Máy vector hỗ trợ",
        [
            "Máy vector hỗ trợ tìm siêu phẳng phân tách hai lớp sao cho lề giữa hai lớp là lớn "
            "nhất. Chỉ những điểm nằm sát lề, gọi là vector hỗ trợ, mới ảnh hưởng tới vị trí "
            "siêu phẳng.",
            "Khi dữ liệu không phân tách tuyến tính được, Kernel Trick cho phép tính tích vô "
            "hướng trong một không gian đặc trưng nhiều chiều hơn mà không cần biến đổi tường "
            "minh từng điểm sang không gian đó, nhờ vậy chi phí tính toán không bùng nổ theo số "
            "chiều.",
        ],
    ),
    (
        "Chương 9 — Tổng kết",
        [
            "Không có thuật toán nào tốt nhất cho mọi bài toán. Việc chọn mô hình luôn phụ thuộc "
            "vào cấu trúc dữ liệu, lượng dữ liệu có sẵn, và ràng buộc về độ trễ khi triển khai.",
            "Một quy trình đánh giá trung thực quan trọng hơn việc chọn được mô hình tinh vi.",
        ],
    ),
]

write_docx("LONG_ML_Handbook.docx", _HANDBOOK_FILLER)

# "Layer Normalization" chỉ xuất hiện ĐÚNG MỘT LẦN, ở section thứ 7/8, trong
# một PDF nhiều trang — kiểm tra đồng thời việc chôn sâu VÀ citation trỏ đúng
# số trang của PDF.
write_pdf(
    "LONG_DL_Handbook.pdf",
    [
        (
            "Chapter 1 - Neurons and Layers",
            [
                "An artificial neuron computes a weighted sum of its inputs, adds a bias, and "
                "passes the result through a non-linear activation function. Stacking neurons "
                "into layers and layers into networks is what gives deep models their capacity.",
                "Without a non-linear activation, any stack of linear layers collapses into a "
                "single linear transformation, no matter how deep it is.",
            ],
        ),
        (
            "Chapter 2 - Activation Functions",
            [
                "ReLU outputs the input when positive and zero otherwise. It is cheap to compute "
                "and does not saturate for positive inputs, which is why it largely replaced "
                "sigmoid and tanh in hidden layers.",
                "A known failure mode of ReLU is the dying unit: if a neuron's input stays "
                "negative, its gradient is always zero and it stops learning. Leaky ReLU "
                "addresses this by allowing a small negative slope.",
            ],
        ),
        (
            "Chapter 3 - Loss Functions",
            [
                "Cross-entropy loss is the standard choice for classification, while mean "
                "squared error is standard for regression. Matching the loss to the task matters "
                "more than small architectural choices.",
                "For imbalanced classification, a weighted loss that increases the cost of "
                "errors on the minority class often works better than resampling the data.",
            ],
        ),
        (
            "Chapter 4 - Optimisers",
            [
                "Stochastic gradient descent updates parameters using the gradient computed on a "
                "mini-batch rather than the whole dataset, trading some gradient accuracy for "
                "far more updates per unit of time.",
                "Adam keeps running averages of both the gradient and its square, which adapts "
                "the effective step size per parameter and usually converges faster than plain "
                "stochastic gradient descent.",
            ],
        ),
        (
            "Chapter 5 - Initialisation",
            [
                "Initialising all weights to the same value makes every neuron in a layer "
                "compute the same thing and receive the same gradient, so the layer never "
                "differentiates. Random initialisation breaks this symmetry.",
                "Xavier initialisation scales the initial weights by the number of input and "
                "output connections, keeping the variance of activations roughly constant across "
                "layers.",
            ],
        ),
        (
            "Chapter 6 - Regularisation in Deep Networks",
            [
                "Dropout randomly disables a fraction of units during training, preventing the "
                "network from relying too heavily on any single unit. At inference time all units "
                "are active and their outputs are scaled accordingly.",
                "Weight decay adds a penalty on the squared magnitude of the weights, which for "
                "plain stochastic gradient descent is equivalent to L2 regularisation.",
            ],
        ),
        (
            "Chapter 7 - Normalisation Inside Networks",
            [
                "Batch normalisation standardises the activations of a layer using statistics "
                "computed across the current mini-batch, which stabilises training and allows "
                "larger learning rates. Its weakness is that it behaves differently at training "
                "and inference time and degrades with very small batches.",
                "Layer Normalization instead standardises across the features of a single "
                "example rather than across the batch, so it behaves identically at training and "
                "inference time and is unaffected by batch size. That property is why it is the "
                "normalisation used inside Transformer blocks.",
            ],
        ),
        (
            "Chapter 8 - Practical Advice",
            [
                "Overfit a tiny subset of the data first. If the model cannot reach near zero "
                "loss on ten examples, there is a bug in the model or the training loop, not a "
                "shortage of data.",
                "Change one thing at a time when tuning. Simultaneous changes make it impossible "
                "to attribute an improvement to any particular decision.",
            ],
        ),
    ],
)

print("\nDone. Corpus now contains:")
for name in sorted(os.listdir(OUT_DIR)):
    print(" ", name)
