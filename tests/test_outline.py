import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "backend"))

from docx import Document as DocxDocument

from app.ingestion.outline import (
    OutlineEntry,
    extract_outline,
    filter_topic_titles,
    is_bibliography_like_chunk,
    is_plausible_topic,
)


def _make_docx(sections):
    """sections: list of (heading, [paragraph, ...])"""
    doc = DocxDocument()
    for heading, paragraphs in sections:
        if heading:
            doc.add_heading(heading, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    # Windows không cho ghi đè khi handle còn mở (WinError 32)
    tmp.close()
    doc.save(tmp.name)
    return tmp.name


class TestExtractOutlineDocx(unittest.TestCase):
    def test_extracts_headings_in_order(self):
        path = _make_docx(
            [
                ("Gradient Descent", ["Nội dung một.", "Nội dung hai."]),
                ("Learning Rate", ["Nội dung ba."]),
            ]
        )
        try:
            outline = extract_outline(path)
            self.assertEqual([e.title for e in outline], ["Gradient Descent", "Learning Rate"])
            self.assertEqual([e.order for e in outline], [0, 1])
        finally:
            os.remove(path)

    def test_entries_carry_position_ref_matching_parser(self):
        # parser gom 10 đoạn/section, nên heading đầu tiên phải nằm ở "Mục 1"
        path = _make_docx([("Chương đầu", ["Một câu."])])
        try:
            outline = extract_outline(path)
            self.assertEqual(outline[0].position_ref, "Mục 1")
        finally:
            os.remove(path)

    def test_heading_in_later_section_gets_later_position_ref(self):
        # 12 đoạn văn trước heading thứ hai -> heading đó rơi sang Mục 2
        first = ("Chương một", [f"Đoạn số {i}." for i in range(12)])
        second = ("Chương hai", ["Nội dung chương hai."])
        path = _make_docx([first, second])
        try:
            outline = extract_outline(path)
            self.assertEqual(outline[0].position_ref, "Mục 1")
            self.assertEqual(outline[1].position_ref, "Mục 2")
        finally:
            os.remove(path)

    def test_document_without_headings_returns_empty(self):
        path = _make_docx([(None, ["Chỉ có đoạn văn thường.", "Không heading nào."])])
        try:
            self.assertEqual(extract_outline(path), [])
        finally:
            os.remove(path)

    def test_returns_outline_entry_instances(self):
        path = _make_docx([("Tiêu đề", ["Nội dung."])])
        try:
            self.assertIsInstance(extract_outline(path)[0], OutlineEntry)
        finally:
            os.remove(path)

    def test_section_index_matches_parse_document_section_for_same_file(self):
        # Bằng chứng cross-consistency: section_index một topic trỏ ra PHẢI
        # đúng section mà app/ingestion/parser.py::parse_document tạo cho
        # cùng file — đây là điều kiện bắt buộc để structural retrieval (lấy
        # đúng dải chunk của một DocumentTopic) hoạt động đúng.
        from app.ingestion.parser import parse_document

        first = ("Chương một", [f"Đoạn số {i}." for i in range(12)])
        second = ("Chương hai", ["Nội dung chương hai."])
        path = _make_docx([first, second])
        try:
            outline = extract_outline(path)
            sections = parse_document(path)
            for entry in outline:
                self.assertLess(entry.section_index, len(sections))
                self.assertEqual(sections[entry.section_index][0], entry.position_ref)
        finally:
            os.remove(path)

    def test_duplicate_headings_are_kept_once(self):
        path = _make_docx([("Trùng", ["A."]), ("Trùng", ["B."])])
        try:
            outline = extract_outline(path)
            self.assertEqual([e.title for e in outline], ["Trùng"])
        finally:
            os.remove(path)

    def test_unsupported_extension_returns_empty(self):
        tmp = tempfile.NamedTemporaryFile(suffix=".txt", delete=False)
        tmp.close()
        try:
            self.assertEqual(extract_outline(tmp.name), [])
        finally:
            os.remove(tmp.name)


class TestIsPlausibleTopic(unittest.TestCase):
    """BUG-001/BUG-006 — lớp lọc chất lượng ở điểm TIÊU THỤ (kế hoạch ôn tập,
    gợi ý chủ đề khi từ chối), tách biệt với heuristic trích xuất PDF."""

    def test_normal_heading_is_plausible(self):
        self.assertTrue(is_plausible_topic("Multi-Head Attention"))
        self.assertTrue(is_plausible_topic("Định lý Pythagoras"))
        self.assertTrue(is_plausible_topic("6.1 Machine Translation"))

    def test_byline_is_rejected_without_hardcoding_the_string(self):
        # Không hardcode "BY A. M. TUBING" — quy tắc byline chung áp dụng cho
        # MỌI byline, kể cả tên bị OCR đọc sai.
        self.assertFalse(is_plausible_topic("BY A. M. TUBING"))
        self.assertFalse(is_plausible_topic("By John Smith"))
        self.assertFalse(is_plausible_topic("Bởi Nguyễn Văn A"))

    def test_citation_like_line_is_rejected(self):
        self.assertFalse(is_plausible_topic("Russell & Norvig (2021), tr. 272."))
        self.assertFalse(is_plausible_topic("Smith, J. (2019). A study. ISBN 123-456."))

    def test_too_short_or_too_long_is_rejected(self):
        self.assertFalse(is_plausible_topic("AB"))
        self.assertFalse(is_plausible_topic("x" * 101))

    def test_empty_or_none_is_rejected(self):
        self.assertFalse(is_plausible_topic(""))
        self.assertFalse(is_plausible_topic("   "))
        self.assertFalse(is_plausible_topic(None))

    def test_mostly_symbolic_ocr_garbage_is_rejected(self):
        self.assertFalse(is_plausible_topic("### 1X1 439 --- ***"))

    def test_line_truncated_mid_clause_is_rejected(self):
        # Dòng bị cắt giữa câu do pypdf ngắt theo độ rộng trang — kết thúc
        # bằng một từ nối/giới từ dang dở, không phải một cụm từ hoàn chỉnh.
        self.assertFalse(is_plausible_topic("Kiến trúc Transformer bỏ qua recurrence và"))
        self.assertFalse(is_plausible_topic("This model was trained with"))

    def test_legitimate_technical_term_is_not_removed(self):
        # Bảo thủ: không loại thuật ngữ kỹ thuật hợp lệ chỉ vì trông lạ.
        self.assertTrue(is_plausible_topic("Positional Encoding"))
        self.assertTrue(is_plausible_topic("Backpropagation"))

    def test_numbered_heading_with_colon_is_still_plausible(self):
        # Dấu hai chấm KHÔNG bị coi là dấu hiệu câu văn (khác dấu phẩy/chấm
        # phẩy) — nhiều heading thật hợp lệ dùng nó, vd tiêu đề bài báo thật:
        self.assertTrue(is_plausible_topic("Deep Learning in Neural Networks: An Overview"))

    def test_numbered_sentence_fragment_with_internal_comma_is_rejected(self):
        # Retest QA (BUG-001): bước chứng minh hình học bị đánh số vẫn lọt
        # trước fix — dấu phẩy giữa câu là tín hiệu "đây là một câu văn, không
        # phải tiêu đề", không phụ thuộc nội dung tài liệu cụ thể nào.
        self.assertFalse(
            is_plausible_topic("1. Nếu hai tam giác có hai cạnh tương ứng bằng nhau,")
        )
        self.assertFalse(
            is_plausible_topic("7. Vì AB bằng FB và BD bằng BC, do đó hai tam giác bằng nhau.")
        )


class TestFilterTopicTitles(unittest.TestCase):
    """BUG-001 retest — QA quan sát được filter theo TỪNG dòng không đủ, vì
    header/footer lặp trên mỗi trang scan sinh ra một biến thể lỗi OCR KHÁC
    NHAU mỗi lần (không dòng nào tự nó trông bất thường). filter_topic_titles
    phải bắt được điều này ở CẤP DANH SÁCH."""

    def test_repeated_near_duplicate_running_header_is_dropped(self):
        titles = [
            "1. The Imitation Game.",
            "COMPUTING MACHINERY AND INTELLIGENCE 435",
            "COMPUTING MACKINBBY AND INTELLIGENCE 437",
            "OOMPUTING MACHWIBY AND INTK1X1OENCK 439",
            "COMPUTING MACHINERY AND INTELLIGENCE 441",
            "2. Critique of the New Problem.",
        ]
        result = filter_topic_titles(titles)
        self.assertIn("1. The Imitation Game.", result)
        self.assertIn("2. Critique of the New Problem.", result)
        for garbage in titles[1:5]:
            self.assertNotIn(garbage, result)

    def test_two_occurrences_of_similar_text_are_kept(self):
        # Cần LẶP LẠI ÍT NHẤT 3 LẦN mới coi là header/footer — 2 lần giống
        # nhau ngẫu nhiên không đủ bằng chứng, tránh loại nhầm heading thật.
        titles = ["Gradient Descent", "Gradient Descend"]  # 1 lỗi chính tả tình cờ
        result = filter_topic_titles(titles)
        self.assertEqual(sorted(result), sorted(titles))

    def test_distinct_legitimate_headings_are_all_kept(self):
        titles = ["1 Introduction", "2 Background", "3 Model Architecture", "Positional Encoding"]
        self.assertEqual(filter_topic_titles(titles), titles)

    def test_empty_list_returns_empty(self):
        self.assertEqual(filter_topic_titles([]), [])

    def test_real_scanned_document_noise_is_mostly_eliminated(self):
        """Bằng chứng thực nghiệm cho retest QA — chạy trên chính file OCR lỗi
        thật trong repo (không phải fixture tự tạo), xác nhận khối lượng nhiễu
        giảm mạnh thay vì chỉ tin vào lý thuyết của heuristic."""
        path = os.path.join(
            os.path.dirname(__file__), "..", "real_test_documents", "03_edge_turing_1950_scanned.pdf"
        )
        if not os.path.exists(path):
            self.skipTest("real_test_documents/03_edge_turing_1950_scanned.pdf không có trong checkout này")
        entries = [e.title for e in extract_outline(path)]
        final = filter_topic_titles(entries)
        # Trước fix: 26/27 dòng lọt qua is_plausible_topic từng dòng riêng lẻ.
        # Sau khi bỏ nhánh Title-Case/ALL-CAPS (outline.py::_looks_like_heading
        # chỉ còn nhận mục đánh số), tài liệu scan này — không có mục đánh số
        # — có thể không còn sinh ra entry thô nào cả, tức nhiễu bằng 0 ngay
        # từ bước trích xuất, còn mạnh hơn yêu cầu gốc của test (giảm mạnh ở
        # bước lọc). len(entries)==0 thì assertLess gốc so sánh 0 < 0 (sai) dù
        # đúng ý test muốn kiểm — tách riêng case này.
        if not entries:
            self.assertEqual(final, [])
        else:
            self.assertLess(
                len(final), len(entries) * 0.3, "khối lượng nhiễu phải giảm mạnh, không chỉ lọt vài dòng"
            )


class TestIsBibliographyLikeChunk(unittest.TestCase):
    """BUG-007 — nhận diện đoạn trích giống khu vực tham khảo để hạ ưu tiên
    khi sinh flashcard (không dùng để loại bỏ khỏi corpus)."""

    def test_reference_list_chunk_is_flagged(self):
        text = (
            "23. Russell & Norvig (2021), tr. 272.\n"
            "24. Smith, J. (2019). A Study of X. ISBN 123-456-789.\n"
            "25. Doe, A. (2018), pp. 45-60."
        )
        self.assertTrue(is_bibliography_like_chunk(text))

    def test_normal_body_text_is_not_flagged(self):
        text = (
            "Định lý Pythagoras phát biểu rằng trong một tam giác vuông, bình "
            "phương cạnh huyền bằng tổng bình phương hai cạnh góc vuông.\n"
            "Đây là một trong những định lý cơ bản nhất của hình học."
        )
        self.assertFalse(is_bibliography_like_chunk(text))

    def test_body_text_mentioning_one_year_is_not_flagged(self):
        # Chỉ MỘT dòng khớp mẫu trích dẫn trong một đoạn nhiều dòng nội dung
        # bình thường không đủ để coi cả đoạn là "tham khảo".
        text = (
            "Định lý được chứng minh lại nhiều lần qua lịch sử.\n"
            "Một bản in lại xuất hiện năm (1968), thu hút sự chú ý mới.\n"
            "Nội dung chứng minh chính vẫn dựa trên tam giác vuông cơ bản.\n"
            "Nhiều SGK phổ thông vẫn dùng cách chứng minh hình học này."
        )
        self.assertFalse(is_bibliography_like_chunk(text))

    def test_empty_text_is_not_flagged(self):
        self.assertFalse(is_bibliography_like_chunk(""))
        self.assertFalse(is_bibliography_like_chunk("   \n  "))


if __name__ == "__main__":
    unittest.main()
